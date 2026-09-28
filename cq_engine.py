import streamlit as st
import google.generativeai as genai
import json
import PyPDF2
from docx import Document

def extract_text_from_pdf(pdf_file):
    reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted: text += extracted + "\n"
    return text

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    text = ""
    for para in doc.paragraphs:
        if para.text: text += para.text + "\n"
    return text

def generate_cqs(context_text, user_topics, custom_instructions, num_easy, num_medium, num_hard, selected_model):
    model = genai.GenerativeModel(selected_model)
    total_questions = num_easy + num_medium + num_hard
    
    if context_text.strip():
        source_instruction = "Based on the provided Context Text and user instructions, generate"
        context_block = f"Context Text:\n{context_text}"
    else:
        source_instruction = "Based on your expert general knowledge and user instructions, generate"
        context_block = "No context text provided. Generate purely based on the requested topics and instructions."
    
    prompt = f"""
    Act as an expert educator. {source_instruction} exactly {total_questions} board-standard Creative Questions (CQ / সৃজনশীল প্রশ্ন).
    
    A standard CQ must include:
    1. "stem": A scenario or context (উদ্দীপক).
    2. "q_a": Knowledge-based question (১ নম্বর).
    3. "q_b": Comprehension-based question (২ নম্বর).
    4. "q_c": Application-based question (৩ নম্বর).
    5. "q_d": Higher-Order Thinking-based question (৪ নম্বর).
    6. Corresponding solutions for each ("sol_a", "sol_b", "sol_c", "sol_d").

    SPECIAL USER GENERATION INSTRUCTIONS (CRITICAL):
    {custom_instructions if custom_instructions else "None provided. Follow standard board-level CQ generation."}

    CRITICAL INSTRUCTIONS FOR DIFFICULTY LEVEL:
    Generate EXACTLY: {num_easy} "Easy", {num_medium} "Medium", {num_hard} "Hard".

    CRITICAL INSTRUCTION FOR THE 'topics' COLUMN FIELD:
    Allowed topic names: {user_topics}
    Select EXACTLY ONE topic from the list. DO NOT include sequence identifiers like '1.0', '2.1'.

    Output STRICTLY as a JSON array of objects with exact keys:
    "sl_no", "stem", "q_a", "q_b", "q_c", "q_d", "sol_a", "sol_b", "sol_c", "sol_d", "subject", "chapter", "topics", "difficulty_level".
    
    {context_block}
    """
    response = model.generate_content(prompt, generation_config=genai.GenerationConfig(response_mime_type="application/json"))
    return json.loads(response.text)

def parse_existing_cqs(raw_cq_text, user_topics, special_instructions, selected_model):
    model = genai.GenerativeModel(selected_model)
    prompt = f"""
    You are an expert educational content parser. Map the provided raw Creative Questions (CQs) strictly into a standard JSON structure.
    Extract the Stem (উদ্দীপক), the 4 questions (A, B, C, D), and their solutions. If solutions are missing, generate accurate ones.
    
    Allowed topics: {user_topics}
    
    SPECIAL USER INSTRUCTIONS (CRITICAL): {special_instructions if special_instructions else "None"}

    Output STRICTLY as a JSON array of objects with keys:
    "sl_no", "stem", "q_a", "q_b", "q_c", "q_d", "sol_a", "sol_b", "sol_c", "sol_d", "subject", "chapter", "topics", "difficulty_level".

    Raw CQs:
    {raw_cq_text}
    """
    response = model.generate_content(prompt, generation_config=genai.GenerationConfig(response_mime_type="application/json"))
    return json.loads(response.text)

def create_cq_docx(cq_data, output_filename="CQs.docx"):
    doc = Document()
    doc.add_heading('Generated CQs', 0)
    table = doc.add_table(rows=1, cols=14)
    table.style = 'Table Grid'
    headers = ["Sl no.", "Stem", "Q (a)", "Q (b)", "Q (c)", "Q (d)", "Sol (a)", "Sol (b)", "Sol (c)", "Sol (d)", "Subject", "Chapter", "Topics", "Difficulty"]
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers): hdr_cells[i].text = header
    for cq in cq_data:
        row_cells = table.add_row().cells
        row_cells[0].text = str(cq.get("sl_no", ""))
        row_cells[1].text = str(cq.get("stem", ""))
        row_cells[2].text = str(cq.get("q_a", ""))
        row_cells[3].text = str(cq.get("q_b", ""))
        row_cells[4].text = str(cq.get("q_c", ""))
        row_cells[5].text = str(cq.get("q_d", ""))
        row_cells[6].text = str(cq.get("sol_a", ""))
        row_cells[7].text = str(cq.get("sol_b", ""))
        row_cells[8].text = str(cq.get("sol_c", ""))
        row_cells[9].text = str(cq.get("sol_d", ""))
        row_cells[10].text = str(cq.get("subject", ""))
        row_cells[11].text = str(cq.get("chapter", ""))
        row_cells[12].text = str(cq.get("topics", ""))
        row_cells[13].text = str(cq.get("difficulty_level", ""))
    doc.save(output_filename)

def run_cq_interface(model_choice, api_key_input):
    # Initialize session state variables specific to CQs
    if "cq_data_t1" not in st.session_state:
        st.session_state["cq_data_t1"] = None
    if "cq_data_t2" not in st.session_state:
        st.session_state["cq_data_t2"] = None
    tab1, tab2 = st.tabs(["✨ Generate New CQs", "📋 Format Existing CQs"])
    
    with tab1:
        with st.container(border=True):
            st.markdown("#### 📄 1. Source Material (Optional)")
            st.caption("Leave blank to generate purely on topics provided.")
            col_file, col_text = st.columns(2)
            with col_file: uploaded_file = st.file_uploader("Upload Document (PDF/DOCX)", type=["pdf", "docx"], key="cq_f1")
            with col_text: raw_text = st.text_area("Or paste raw text here", height=100, key="cq_t1")

        with st.container(border=True):
            st.markdown("#### 🎯 2. Question Parameters")
            topics_input = st.text_input("Allowed Topics for Table Mapping (comma-separated)", placeholder="e.g., Computer Network", key="cq_top1")
            custom_instructions = st.text_area("Custom Generation Instructions (Optional)", height=100, key="cq_cust1")

        with st.container(border=True):
            st.markdown("#### 📊 3. Difficulty Breakdown")
            col1, col2, col3 = st.columns(3)
            with col1: num_easy = st.number_input("🟢 Easy", min_value=0, value=2, key="cq_e1")
            with col2: num_medium = st.number_input("🟡 Medium", min_value=0, value=2, key="cq_m1")
            with col3: num_hard = st.number_input("🔴 Hard", min_value=0, value=1, key="cq_h1")

        total_q = num_easy + num_medium + num_hard
        st.markdown("<br>", unsafe_allow_html=True)
        
        col_btn, col_msg = st.columns([1, 2])
        with col_btn: generate_btn = st.button("🚀 Generate CQs", use_container_width=True, type="primary")
        with col_msg: st.info(f"Target Generation: **{total_q}** Creative Questions")

        if generate_btn:
            if not api_key_input: st.error("Please enter your API Key in the sidebar first!")
            elif total_q == 0: st.error("Please specify at least one question to generate.")
            elif not topics_input: st.error("Please provide at least one topic.")
            else:
                with st.spinner(f"Engine running... Generating {total_q} CQs using {model_choice}"):
                    text_to_process = ""
                    if uploaded_file:
                        if uploaded_file.name.endswith('.pdf'): text_to_process = extract_text_from_pdf(uploaded_file)
                        elif uploaded_file.name.endswith('.docx'): text_to_process = extract_text_from_docx(uploaded_file)
                    elif raw_text: text_to_process = raw_text
                    try:
                        st.session_state["cq_data_t1"] = generate_cqs(text_to_process, topics_input, custom_instructions, num_easy, num_medium, num_hard, model_choice)
                        st.success("✨ Generation Complete!")
                    except Exception as e: st.error(f"An error occurred: {e}")

        if st.session_state["cq_data_t1"]:
            with st.container(border=True):
                st.markdown("#### 👀 Preview & Edit Generated Table")
                edited_data_cq1 = st.data_editor(st.session_state["cq_data_t1"], use_container_width=True, num_rows="dynamic", key="cq_edit1")
            create_cq_docx(edited_data_cq1, "Generated_CQs.docx")
            with open("Generated_CQs.docx", "rb") as file:
                st.download_button("📥 Download Edited Word Document (.docx)", data=file, file_name="Generated_CQs.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)

    with tab2:
        with st.container(border=True):
            st.markdown("#### 📝 1. Raw Input")
            col_file_t2, col_text_t2 = st.columns(2)
            with col_file_t2: uploaded_file_t2 = st.file_uploader("Upload Raw CQs (PDF/DOCX)", type=["pdf", "docx"], key="cq_f2")
            with col_text_t2: raw_text_t2 = st.text_area("Or paste raw pre-written CQs here", height=150, key="cq_t2")
        
        with st.container(border=True):
            st.markdown("#### 🎯 2. Formatting Parameters")
            topics_input_t2 = st.text_input("Target Topics (Optional)", key="cq_top2")
            special_instructions = st.text_area("Special AI Instructions (Optional)", height=100, key="cq_cust2")

        st.markdown("<br>", unsafe_allow_html=True)
        if st.button("🛠️ Format Existing CQs", use_container_width=True, type="primary"):
            if not api_key_input: st.error("Please configure your API Key in the sidebar.")
            elif not uploaded_file_t2 and not raw_text_t2: st.error("Please provide raw CQs.")
            else:
                with st.spinner("Parsing and reformatting CQs..."):
                    text_to_process = ""
                    if uploaded_file_t2:
                        if uploaded_file_t2.name.endswith('.pdf'): text_to_process = extract_text_from_pdf(uploaded_file_t2)
                        elif uploaded_file_t2.name.endswith('.docx'): text_to_process = extract_text_from_docx(uploaded_file_t2)
                    elif raw_text_t2: text_to_process = raw_text_t2
                    try:
                        st.session_state["cq_data_t2"] = parse_existing_cqs(text_to_process, topics_input_t2, special_instructions, model_choice)
                        st.success("✨ Successfully reformatted!")
                    except Exception as e: st.error(f"Error parsing CQs: {e}")

        if st.session_state["cq_data_t2"]:
            with st.container(border=True):
                st.markdown("#### 👀 Preview & Edit Formatted Table")
                edited_data_cq2 = st.data_editor(st.session_state["cq_data_t2"], use_container_width=True, num_rows="dynamic", key="cq_edit2")
            create_cq_docx(edited_data_cq2, "Formatted_Ready_CQs.docx")
            with open("Formatted_Ready_CQs.docx", "rb") as f:
                st.download_button("📥 Download Edited Word Document (.docx)", data=f, file_name="Formatted_Ready_CQs.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
