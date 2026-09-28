import streamlit as st
import google.generativeai as genai
import json
import PyPDF2
from docx import Document

# --- 1. SECURE API KEY INPUT ---
st.sidebar.title("⚙️ Configuration")
api_key = st.sidebar.text_input("Enter your Gemini API Key:", type="password")

if api_key:
    genai.configure(api_key=api_key)
else:
    st.sidebar.warning("Please enter your API key to proceed.") 

# --- 2. HELPER FUNCTIONS ---
def extract_text_from_pdf(pdf_file):
    """Reads the uploaded PDF and pulls all the text out."""
    reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            text += extracted + "\n"
    return text

def generate_mcqs(context_text, user_topics, num_easy, num_medium, num_hard):
    """Sends the text, topics, and difficulty distribution using the legacy SDK."""
    model = genai.GenerativeModel('gemini-3.1-flash-lite')
    total_questions = num_easy + num_medium + num_hard
    
    prompt = f"""
    Act as an expert educator. Based on the following Context Text, generate exactly {total_questions} board-standard multiple-choice questions.
    
    CRITICAL INSTRUCTIONS FOR DIFFICULTY LEVEL:
    You must generate EXACTLY:
    - {num_easy} questions where "difficulty_level" is "Easy"
    - {num_medium} questions where "difficulty_level" is "Medium"
    - {num_hard} questions where "difficulty_level" is "Hard"
    Do not use any other words for difficulty level.

    CRITICAL INSTRUCTION FOR THE 'topics' FIELD:
    Here is a list of allowed topics: {user_topics}
    For the "topics" field in each question, you MUST select EXACTLY ONE topic from the list above that best fits the question. Do not invent any new topics.

    Output the result STRICTLY as a JSON array of objects. Do not include markdown formatting like ```json.
    Each object must have the following exact keys:
    "sl_no", "question_title", "A", "B", "C", "D", 
    "solution_body", "correct_option", "subject", "chapter", "topics", 
    "question_category", "difficulty_level".
    
    Context Text:
    {context_text}
    """
    
    response = model.generate_content(
        prompt,
        generation_config=genai.GenerationConfig(
            response_mime_type="application/json",
        )
    )
    
    return json.loads(response.text)

def create_mcq_docx(mcq_data, output_filename="MCQs.docx"):
    """Creates the formatted Word Document with the 13-column table."""
    doc = Document()
    doc.add_heading('Generated MCQs', 0)
    
    table = doc.add_table(rows=1, cols=13)
    table.style = 'Table Grid'
    
    headers = [
        "Sl no.", "Question Title", "Option A", "Option B", "Option C", 
        "Option D", "Solution Body", "Correct Option", "Subject", 
        "Chapter", "Topics", "Question Category", "Difficulty Level"
    ]
    
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        
    for mcq in mcq_data:
        row_cells = table.add_row().cells
        row_cells[0].text = str(mcq.get("sl_no", ""))
        row_cells[1].text = str(mcq.get("question_title", ""))
        row_cells[2].text = str(mcq.get("A", ""))
        row_cells[3].text = str(mcq.get("B", ""))
        row_cells[4].text = str(mcq.get("C", ""))
        row_cells[5].text = str(mcq.get("D", ""))
        row_cells[6].text = str(mcq.get("solution_body", ""))
        row_cells[7].text = str(mcq.get("correct_option", ""))
        row_cells[8].text = str(mcq.get("subject", ""))
        row_cells[9].text = str(mcq.get("chapter", ""))
        row_cells[10].text = str(mcq.get("topics", ""))
        row_cells[11].text = str(mcq.get("question_category", ""))
        row_cells[12].text = str(mcq.get("difficulty_level", ""))
        
    doc.save(output_filename)

# --- 3. WEB INTERFACE (STREAMLIT) ---
st.title("Automated MCQ Generator")
st.write("Upload a source document, define your topics, and configure difficulty levels.")

uploaded_file = st.file_uploader("1. Upload Source PDF", type="pdf")
raw_text = st.text_area("Or paste raw source text here")
topics_input = st.text_input("2. Enter specific topics (comma-separated)", placeholder="e.g., Hardware, Memory, Super Computers")

st.write("3. Difficulty Breakdown (Number of Questions)")
col1, col2, col3 = st.columns(3)
with col1:
    num_easy = st.number_input("Easy", min_value=0, max_value=50, value=10)
with col2:
    num_medium = st.number_input("Medium", min_value=0, max_value=50, value=15)
with col3:
    num_hard = st.number_input("Hard", min_value=0, max_value=50, value=5)

total_q = num_easy + num_medium + num_hard
st.info(f"Total questions to generate: {total_q}")

if st.button("Generate MCQs"):
    if not api_key:
        st.error("Please enter your API Key in the sidebar first!")
    if total_q == 0:
        st.error("Please specify at least one question to generate.")
    elif not topics_input:
        st.error("Please provide at least one topic.")
    elif not uploaded_file and not raw_text:
        st.error("Please provide some source text or upload a PDF.")
    else:
        with st.spinner(f"Analyzing text and generating {total_q} questions (this takes a minute)..."):
            text_to_process = ""
            if uploaded_file:
                text_to_process = extract_text_from_pdf(uploaded_file)
            elif raw_text:
                text_to_process = raw_text
                
            try:
                mcq_json = generate_mcqs(text_to_process, topics_input, num_easy, num_medium, num_hard)
                create_mcq_docx(mcq_json, "Generated_MCQs.docx")
                
                st.success("Successfully generated!")
                
                with open("Generated_MCQs.docx", "rb") as file:
                    st.download_button(
                        label="Download Formatted Word Document",
                        data=file,
                        file_name="Generated_MCQs.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
            except Exception as e:
                st.error(f"An error occurred: {e}")